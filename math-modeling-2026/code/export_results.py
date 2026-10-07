"""Export computed fields to the four result workbooks and complete CSV tables."""
from pathlib import Path
import csv,json
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment

PACKAGE=Path(__file__).resolve().parents[1]

def table(a,field,moving):
    rr=np.arange(21)/1000
    header=['时间/s']+[f'{v*100:.1f} cm' for v in rr]
    if moving:header+=['药材表面','半径/cm']
    data=[]
    for t,r,y in zip(a['history'][:,0],a['r_m'],a[field]):
        value=np.interp(rr,r,y).astype(object);value[rr>r[-1]+1e-12]=None
        row=[float(t)]+[None if v is None else round(float(v),4) for v in value]
        if moving:row+=[round(float(y[-1]),4),round(float(r[-1]*100),4)]
        data.append(row)
    return header,data

def main():
    summary=[]
    for number,tag,fields in [(1,'q1',['T_C','C']),(2,'q2',['T_C','C']),(3,'q3',['C']),(4,'q4_uniform',['C'])]:
        a=np.load(PACKAGE/'results/raw'/f'{tag}.npz');wb=Workbook();wb.remove(wb.active)
        wb.properties.creator='';wb.properties.lastModifiedBy='';wb.properties.title=f'问题{number}复算结果'
        for field in fields:
            title='温度' if field=='T_C' else '水分浓度';header,rows=table(a,field,number==4)
            ws=wb.create_sheet(title);ws.append(header)
            for row in rows:ws.append(row)
            ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='234F70');cell.alignment=Alignment(horizontal='center')
            for column in ws.columns:
                ws.column_dimensions[column[0].column_letter].width=13
                for cell in column[1:]:cell.number_format='0.0000' if cell.column>1 else '0.####'
            with (PACKAGE/'results'/f'result{number}_{field}.csv').open('w',encoding='utf-8-sig',newline='') as f:
                w=csv.writer(f);w.writerow(header);w.writerows(rows)
        wb.save(PACKAGE/'results'/f'result{number}.xlsx')
        j=json.loads((PACKAGE/'results/raw'/f'{tag}.json').read_text())
        summary.append({'problem':number,'event_h':j['event_h'],'N':j['N'],'dt_s':j['dt_s'],'water_balance':j['diagnostics']['final_relative_water_balance']})
    a=np.load(PACKAGE/'results/raw/q4_mass.npz');header,rows=table(a,'C',True)
    with (PACKAGE/'results/Q4_mass_complete.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(header);w.writerows(rows)
    (PACKAGE/'results/summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
