"""Read and validate the official input workbooks."""
from pathlib import Path
import hashlib,json
import numpy as np
from openpyxl import load_workbook

PACKAGE=Path(__file__).resolve().parents[2]
EXPECTED={'附件1.xlsx':'7ef32870abeef420b89560b2530ff60dfe4255917805151d89988d0311af9dd7',
          '附件2.xlsx':'5563acbfa4b4afb10cc6c03e2207e5369bf39da27576672aff14cc5c32e704af'}

def prepare(folder):
    folder=Path(folder);dest=PACKAGE/'data/processed';dest.mkdir(parents=True,exist_ok=True)
    (PACKAGE/'results/raw').mkdir(parents=True,exist_ok=True)
    records=[]
    for idx,cols,rows in [(1,3,241),(2,2,145)]:
        name=f'附件{idx}.xlsx';p=folder/name
        if not p.is_file():raise FileNotFoundError(f'Official input missing: {name}; use --data-dir.')
        digest=hashlib.sha256(p.read_bytes()).hexdigest()
        if digest!=EXPECTED[name]:raise ValueError(f'Input differs from the reviewed official file: {name}')
        wb=load_workbook(p,read_only=True,data_only=True)
        values=[r[:cols] for r in wb.worksheets[0].iter_rows(values_only=True)
                if len(r)>=cols and all(isinstance(v,(int,float)) for v in r[:cols])]
        wb.close();a=np.asarray(values,float)
        assert a.shape==(rows,cols) and np.all(np.diff(a[:,0])>0)
        np.save(dest/f'attachment{idx}.npy',a)
        records.append({'file':name,'sha256':digest,'rows':rows})
    (PACKAGE/'results/input_validation.json').write_text(json.dumps(records,ensure_ascii=False,indent=2))
